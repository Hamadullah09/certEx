"""Downloading a batch.

This is the moment certificate details leave the system, so two things are asserted
throughout: the file contains exactly the rows the filters asked for, and the fact that
it happened is written down with who did it.
"""

from __future__ import annotations

import csv
import io
import json
import uuid
import zipfile

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from certex.db.models import AuditLog, Batch, CertificateUnit, Document, Export, Extraction, User
from certex.enums import (
    AuditAction,
    CertificateType,
    DocumentStatus,
    ExportFormat,
    ExtractionMethod,
    ReviewStatus,
    UnitStatus,
)
from certex.export.writers import UTF8_BOM
from tests.conftest import authenticate

pytestmark = pytest.mark.integration

BIRTH = {
    "certificate_number": "BC-2019-004471",
    "child_full_name": "Ayesha Noor Malik",
    "date_of_birth": "1987-03-14",
    "sex": "F",
}
DEATH = {
    "certificate_number": "DC-2021-000913",
    "deceased_full_name": "Abdul Rehman Qureshi",
    "date_of_death": "2021-11-17",
    "cause_of_death": "Cardiac arrest",
}


async def add_row(
    session: AsyncSession,
    batch: Batch,
    *,
    fields: dict[str, str],
    certificate_type: CertificateType = CertificateType.BIRTH,
    review_status: ReviewStatus = ReviewStatus.NEEDS_REVIEW,
    flags: list[str] | None = None,
    extras: dict[str, str] | None = None,
    file_name: str = "birth.pdf",
) -> Extraction:
    document = Document(
        batch_id=batch.id,
        workspace_id=batch.workspace_id,
        original_filename=file_name,
        mime_type="application/pdf",
        byte_size=1024,
        sha256=uuid.uuid4().hex * 2,
        storage_key=f"documents/{batch.workspace_id}/2026/09/{uuid.uuid4()}",
        status=DocumentStatus.COMPLETED,
        page_count=2,
    )
    session.add(document)
    await session.flush()
    unit = CertificateUnit(
        document_id=document.id,
        batch_id=batch.id,
        ordinal=0,
        page_start=1,
        page_end=2,
        certificate_type=certificate_type,
        status=UnitStatus.COMPLETED,
    )
    session.add(unit)
    await session.flush()
    row = Extraction(
        unit_id=unit.id,
        batch_id=batch.id,
        workspace_id=batch.workspace_id,
        certificate_type=certificate_type,
        fields_jsonb=dict(fields),
        field_confidences_jsonb=dict.fromkeys(fields, 0.93),
        field_methods_jsonb=dict.fromkeys(fields, ExtractionMethod.RULE.value),
        field_sources_jsonb={
            name: {"page_number": 1, "snippet": value, "bbox": None, "label": name}
            for name, value in fields.items()
        },
        extra_fields_jsonb=dict(extras or {}),
        flags_jsonb=flags or [],
        row_confidence=0.91,
        review_status=review_status,
    )
    session.add(row)
    await session.flush()
    return row


def rows_of(text: str, delimiter: str = ",") -> list[list[str]]:
    return list(csv.reader(io.StringIO(text.lstrip(UTF8_BOM)), delimiter=delimiter))


class TestCsvDownload:
    async def test_the_batch_comes_back_as_a_csv(
        self, api_client: AsyncClient, db_session: AsyncSession, operator_user: User, batch: Batch
    ) -> None:
        await add_row(db_session, batch, fields=BIRTH)
        await authenticate(api_client, operator_user)

        response = await api_client.get(f"/api/v1/batches/{batch.id}/export")

        assert response.status_code == 200, response.text
        assert response.headers["content-type"].startswith("text/csv")
        rows = rows_of(response.text)
        assert rows[0][0] == "Serial no"
        assert rows[1][0] == "1"
        assert "Ayesha Noor Malik" in rows[1]

    async def test_the_file_is_named_after_the_batch_and_the_moment(
        self, api_client: AsyncClient, db_session: AsyncSession, operator_user: User, batch: Batch
    ) -> None:
        await add_row(db_session, batch, fields=BIRTH)
        await authenticate(api_client, operator_user)

        response = await api_client.get(f"/api/v1/batches/{batch.id}/export")

        disposition = response.headers["content-disposition"]
        assert "attachment" in disposition
        assert "certificates_Test-Batch_" in disposition
        assert disposition.endswith(".csv") or ".csv" in disposition

    async def test_rows_come_out_in_serial_order(
        self, api_client: AsyncClient, db_session: AsyncSession, operator_user: User, batch: Batch
    ) -> None:
        for index in range(3):
            await add_row(db_session, batch, fields={**BIRTH, "certificate_number": f"BC-{index}"})
        await authenticate(api_client, operator_user)

        response = await api_client.get(f"/api/v1/batches/{batch.id}/export")

        rows = rows_of(response.text)[1:]
        assert [row[0] for row in rows] == ["1", "2", "3"]

    async def test_a_mixed_batch_gets_one_sheet_with_both_sets_of_columns(
        self, api_client: AsyncClient, db_session: AsyncSession, operator_user: User, batch: Batch
    ) -> None:
        await add_row(db_session, batch, fields=BIRTH)
        await add_row(db_session, batch, fields=DEATH, certificate_type=CertificateType.DEATH)
        await authenticate(api_client, operator_user)

        response = await api_client.get(f"/api/v1/batches/{batch.id}/export")

        header = rows_of(response.text)[0]
        assert "Name of child" in header
        assert "Cause of death" in header

    async def test_a_label_the_schema_never_knew_still_reaches_the_file(
        self, api_client: AsyncClient, db_session: AsyncSession, operator_user: User, batch: Batch
    ) -> None:
        await add_row(db_session, batch, fields=BIRTH, extras={"blood_group": "O+"})
        await authenticate(api_client, operator_user)

        response = await api_client.get(f"/api/v1/batches/{batch.id}/export")

        rows = rows_of(response.text)
        assert any("extra" in header.lower() for header in rows[0])
        assert "O+" in rows[1]

    async def test_where_each_row_came_from_is_in_the_file(
        self, api_client: AsyncClient, db_session: AsyncSession, operator_user: User, batch: Batch
    ) -> None:
        await add_row(db_session, batch, fields=BIRTH, file_name="register-march.pdf")
        await authenticate(api_client, operator_user)

        response = await api_client.get(f"/api/v1/batches/{batch.id}/export")

        rows = rows_of(response.text)
        assert "register-march.pdf" in rows[1]
        assert "1-2" in rows[1], "the page range this row covers"

    async def test_the_delimiter_can_be_chosen(
        self, api_client: AsyncClient, db_session: AsyncSession, operator_user: User, batch: Batch
    ) -> None:
        await add_row(db_session, batch, fields=BIRTH)
        await authenticate(api_client, operator_user)

        response = await api_client.get(f"/api/v1/batches/{batch.id}/export?delimiter=semicolon")

        rows = rows_of(response.text, delimiter=";")
        assert rows[1][0] == "1"

    async def test_confidence_can_be_asked_for(
        self, api_client: AsyncClient, db_session: AsyncSession, operator_user: User, batch: Batch
    ) -> None:
        await add_row(db_session, batch, fields=BIRTH)
        await authenticate(api_client, operator_user)

        response = await api_client.get(
            f"/api/v1/batches/{batch.id}/export?include_confidence=true&include_snippet=true"
        )

        header = rows_of(response.text)[0]
        assert any("(confidence)" in cell for cell in header)
        assert any("(as printed)" in cell for cell in header)


class TestFilters:
    async def test_the_download_honours_the_filter_on_screen(
        self, api_client: AsyncClient, db_session: AsyncSession, operator_user: User, batch: Batch
    ) -> None:
        # A reviewer who filtered to "needs review" expects that file, not everything.
        await add_row(db_session, batch, fields=BIRTH, review_status=ReviewStatus.MANUALLY_APPROVED)
        await add_row(
            db_session,
            batch,
            fields={**BIRTH, "certificate_number": "BC-2"},
            review_status=ReviewStatus.NEEDS_REVIEW,
        )
        await authenticate(api_client, operator_user)

        response = await api_client.get(
            f"/api/v1/batches/{batch.id}/export?filter[status]=NEEDS_REVIEW"
        )

        rows = rows_of(response.text)
        assert len(rows) == 2  # the header and one row
        assert "BC-2" in rows[1]

    async def test_a_type_can_be_downloaded_on_its_own(
        self, api_client: AsyncClient, db_session: AsyncSession, operator_user: User, batch: Batch
    ) -> None:
        await add_row(db_session, batch, fields=BIRTH)
        await add_row(db_session, batch, fields=DEATH, certificate_type=CertificateType.DEATH)
        await authenticate(api_client, operator_user)

        response = await api_client.get(f"/api/v1/batches/{batch.id}/export?filter[type]=DEATH")

        rows = rows_of(response.text)
        assert len(rows) == 2
        assert "Abdul Rehman Qureshi" in rows[1]


class TestOtherFormats:
    async def test_excel(
        self, api_client: AsyncClient, db_session: AsyncSession, operator_user: User, batch: Batch
    ) -> None:
        await add_row(db_session, batch, fields=BIRTH)
        await authenticate(api_client, operator_user)

        response = await api_client.get(f"/api/v1/batches/{batch.id}/export?format=xlsx")

        assert response.status_code == 200
        sheet = load_workbook(io.BytesIO(response.content)).active
        assert sheet is not None
        assert sheet["A1"].value == "Serial no"
        assert sheet["A2"].value == "1"

    async def test_json(
        self, api_client: AsyncClient, db_session: AsyncSession, operator_user: User, batch: Batch
    ) -> None:
        await add_row(db_session, batch, fields=BIRTH)
        await authenticate(api_client, operator_user)

        response = await api_client.get(f"/api/v1/batches/{batch.id}/export?format=json")

        payload = json.loads(response.text)
        assert payload[0]["serial_no"] == "1"
        assert payload[0]["child_full_name"] == "Ayesha Noor Malik"

    async def test_one_file_per_type_as_a_zip(
        self, api_client: AsyncClient, db_session: AsyncSession, operator_user: User, batch: Batch
    ) -> None:
        await add_row(db_session, batch, fields=BIRTH)
        await add_row(db_session, batch, fields=DEATH, certificate_type=CertificateType.DEATH)
        await authenticate(api_client, operator_user)

        response = await api_client.get(f"/api/v1/batches/{batch.id}/export?per_type=true")

        assert response.headers["content-type"] == "application/zip"
        with zipfile.ZipFile(io.BytesIO(response.content)) as bundle:
            names = bundle.namelist()
            assert len(names) == 2
            assert any("birth" in name for name in names)
            assert any("death" in name for name in names)
            birth_file = next(name for name in names if "birth" in name)
            text = bundle.read(birth_file).decode("utf-8")
            assert "Ayesha Noor Malik" in text
            assert "Abdul Rehman Qureshi" not in text


class TestPreview:
    async def test_it_says_what_the_download_would_contain(
        self, api_client: AsyncClient, db_session: AsyncSession, operator_user: User, batch: Batch
    ) -> None:
        await add_row(db_session, batch, fields=BIRTH)
        await add_row(db_session, batch, fields=DEATH, certificate_type=CertificateType.DEATH)
        await authenticate(api_client, operator_user)

        response = await api_client.get(f"/api/v1/batches/{batch.id}/export/preview")

        body = response.json()
        assert body["row_count"] == 2
        assert body["column_count"] > 10
        assert set(body["certificate_types"]) == {"BIRTH", "DEATH"}
        assert body["columns"][0] == "Serial no"
        assert body["filename"].startswith("certificates_")

    async def test_it_warns_how_many_rows_nobody_has_checked(
        self, api_client: AsyncClient, db_session: AsyncSession, operator_user: User, batch: Batch
    ) -> None:
        # Exporting unchecked rows is allowed; being surprised by it is not.
        await add_row(db_session, batch, fields=BIRTH, review_status=ReviewStatus.MANUALLY_APPROVED)
        await add_row(
            db_session,
            batch,
            fields={**BIRTH, "certificate_number": "BC-2"},
            review_status=ReviewStatus.NEEDS_REVIEW,
        )
        await add_row(
            db_session,
            batch,
            fields={**BIRTH, "certificate_number": "BC-3"},
            review_status=ReviewStatus.FAILED,
        )
        await authenticate(api_client, operator_user)

        response = await api_client.get(f"/api/v1/batches/{batch.id}/export/preview")

        assert response.json()["unreviewed_count"] == 2

    async def test_the_preview_matches_the_filters(
        self, api_client: AsyncClient, db_session: AsyncSession, operator_user: User, batch: Batch
    ) -> None:
        await add_row(db_session, batch, fields=BIRTH)
        await add_row(db_session, batch, fields=DEATH, certificate_type=CertificateType.DEATH)
        await authenticate(api_client, operator_user)

        response = await api_client.get(
            f"/api/v1/batches/{batch.id}/export/preview?filter[type]=BIRTH"
        )

        assert response.json()["row_count"] == 1


class TestRecordKeeping:
    async def test_every_download_is_recorded(
        self, api_client: AsyncClient, db_session: AsyncSession, operator_user: User, batch: Batch
    ) -> None:
        await add_row(db_session, batch, fields=BIRTH)
        await authenticate(api_client, operator_user)

        await api_client.get(f"/api/v1/batches/{batch.id}/export")

        export = await db_session.scalar(select(Export).where(Export.batch_id == batch.id))
        assert export is not None
        assert export.format is ExportFormat.CSV
        assert export.row_count == 1
        assert export.created_by == operator_user.id

    async def test_the_download_is_audited(
        self, api_client: AsyncClient, db_session: AsyncSession, operator_user: User, batch: Batch
    ) -> None:
        await add_row(db_session, batch, fields=BIRTH)
        await authenticate(api_client, operator_user)

        await api_client.get(f"/api/v1/batches/{batch.id}/export?format=json")

        entry = await db_session.scalar(
            select(AuditLog).where(AuditLog.action == AuditAction.EXPORT_REQUESTED)
        )
        assert entry is not None
        assert entry.entity_id == batch.id
        assert entry.metadata_jsonb["format"] == "json"


class TestAccess:
    async def test_a_viewer_may_download(
        self, api_client: AsyncClient, db_session: AsyncSession, viewer_user: User, batch: Batch
    ) -> None:
        # Reading the data is what a viewer is for.
        await add_row(db_session, batch, fields=BIRTH)
        await authenticate(api_client, viewer_user)

        response = await api_client.get(f"/api/v1/batches/{batch.id}/export")

        assert response.status_code == 200

    async def test_another_workspaces_batch_cannot_be_downloaded(
        self, api_client: AsyncClient, operator_user: User
    ) -> None:
        await authenticate(api_client, operator_user)
        response = await api_client.get(f"/api/v1/batches/{uuid.uuid4()}/export")
        assert response.status_code == 404

    async def test_signing_in_is_required(self, api_client: AsyncClient, batch: Batch) -> None:
        response = await api_client.get(f"/api/v1/batches/{batch.id}/export")
        assert response.status_code == 401


class TestEmptyBatch:
    async def test_a_batch_with_no_rows_still_downloads_a_header(
        self, api_client: AsyncClient, operator_user: User, batch: Batch
    ) -> None:
        # An empty file with the right columns is honest; an error is confusing.
        await authenticate(api_client, operator_user)

        response = await api_client.get(f"/api/v1/batches/{batch.id}/export")

        rows = rows_of(response.text)
        assert len(rows) == 1
        assert rows[0][0] == "Serial no"
