"""Writing the file an office actually opens.

The tests that matter here are not about formatting. They are about a file that arrives
on a clerk's Windows machine, is opened in Excel, and must not execute anything, must not
turn Urdu names into mojibake, and must put the serial number in the first column where
everyone expects it.
"""

from __future__ import annotations

import csv
import io
import json

import pytest
from openpyxl import load_workbook

from certex.enums import CertificateType
from certex.export.columns import EXTRA_PREFIX, build_column_plan
from certex.export.rows import ExportRow
from certex.export.writers import (
    DELIMITERS,
    UTF8_BOM,
    guard_formula,
    write_csv,
    write_json,
    write_xlsx,
)

pytestmark = pytest.mark.unit


def plan(*types: CertificateType, **kwargs: object):
    return build_column_plan(types or (CertificateType.BIRTH,), **kwargs)  # type: ignore[arg-type]


def row(**cells: str) -> ExportRow:
    flagged = bool(cells.pop("_flagged", ""))
    return ExportRow(cells=dict(cells), flagged=flagged)


def csv_text(columns, rows, **kwargs: object) -> str:
    return "".join(write_csv(columns, rows, **kwargs))  # type: ignore[arg-type]


def parsed(text: str, delimiter: str = ",") -> list[list[str]]:
    return list(csv.reader(io.StringIO(text.lstrip(UTF8_BOM)), delimiter=delimiter))


class TestColumnOrder:
    def test_the_serial_number_comes_first_and_the_type_second(self) -> None:
        # This is how an office refers to a row, so it is the first thing they see.
        headers = plan().headers
        assert headers[0] == "Serial no"
        assert headers[1] == "Certificate type"

    def test_the_fields_every_certificate_has_come_before_the_type_specific_ones(self) -> None:
        keys = plan().keys
        assert keys.index("certificate_number") < keys.index("child_full_name")

    def test_a_mixed_batch_gets_one_sheet_with_every_type_s_fields(self) -> None:
        keys = plan(CertificateType.BIRTH, CertificateType.DEATH).keys
        assert "child_full_name" in keys
        assert "cause_of_death" in keys

    def test_the_columns_do_not_depend_on_the_order_the_types_arrived_in(self) -> None:
        # Two exports of one batch must have identical columns.
        first = plan(CertificateType.BIRTH, CertificateType.DEATH).keys
        second = plan(CertificateType.DEATH, CertificateType.BIRTH).keys
        assert first == second

    def test_unknown_labels_are_marked_and_sorted(self) -> None:
        keys = plan(extra_field_names=["blood_group", "annexure"]).keys
        assert keys.index(f"{EXTRA_PREFIX}annexure") < keys.index(f"{EXTRA_PREFIX}blood_group")

    def test_where_a_row_came_from_is_always_in_the_file(self) -> None:
        keys = plan().keys
        for name in ("file_name", "page_range", "review_status", "row_confidence", "flags"):
            assert name in keys

    def test_confidence_and_source_are_off_unless_asked_for(self) -> None:
        assert not any("confidence)" in header for header in plan().headers)
        with_confidence = plan(include_confidence=True).headers
        assert any("(confidence)" in header for header in with_confidence)
        with_source = plan(include_snippet=True).headers
        assert any("(as printed)" in header for header in with_source)


class TestFormulaInjection:
    @pytest.mark.parametrize("value", ["=1+1", "+44 300", "-2019-004471", "@SUM(A1)"])
    def test_a_cell_a_spreadsheet_would_execute_is_neutralised(self, value: str) -> None:
        # The office that opens the file is the one that would pay for this.
        assert guard_formula(value).startswith("'")

    @pytest.mark.parametrize("value", ["BC-2019-004471", "Ayesha Noor Malik", "", "35201-1"])
    def test_ordinary_values_are_untouched(self, value: str) -> None:
        assert guard_formula(value) == value

    def test_the_guard_reaches_the_written_file(self) -> None:
        columns = plan()
        text = csv_text(columns, [row(serial_no="1", certificate_number="=cmd|'/c calc'!A1")])
        assert "'=cmd" in text

    def test_headers_are_guarded_too(self) -> None:
        assert guard_formula("=Name") == "'=Name"


class TestCsv:
    def test_the_header_names_every_column(self) -> None:
        columns = plan()
        rows = parsed(csv_text(columns, []))
        assert rows[0] == columns.headers

    def test_a_value_that_was_not_found_is_an_empty_cell(self) -> None:
        # Not "None", not "null", not a dash: empty means nothing was there.
        columns = plan()
        rows = parsed(csv_text(columns, [row(serial_no="1")]))
        assert set(rows[1][2:]) <= {""}

    def test_excel_on_windows_reads_urdu_correctly(self) -> None:
        columns = plan()
        text = csv_text(columns, [row(serial_no="1", child_full_name="عائشہ نور ملک")])
        assert text.startswith(UTF8_BOM), "without the mark, Excel mangles every Urdu name"
        assert "عائشہ نور ملک" in text

    def test_the_mark_can_be_left_off_for_everything_that_is_not_excel(self) -> None:
        text = csv_text(plan(), [], include_bom=False)
        assert not text.startswith(UTF8_BOM)

    @pytest.mark.parametrize("name", list(DELIMITERS))
    def test_every_offered_delimiter_round_trips(self, name: str) -> None:
        delimiter = DELIMITERS[name]
        columns = plan()
        text = csv_text(
            columns, [row(serial_no="1", child_full_name="Ayesha")], delimiter=delimiter
        )
        rows = parsed(text, delimiter=delimiter)
        assert rows[1][0] == "1"

    def test_a_value_containing_the_delimiter_survives(self) -> None:
        columns = plan()
        text = csv_text(columns, [row(serial_no="1", place_of_birth="Services Hospital, Lahore")])
        rows = parsed(text)
        assert "Services Hospital, Lahore" in rows[1]

    def test_a_value_containing_a_newline_survives(self) -> None:
        columns = plan()
        text = csv_text(columns, [row(serial_no="1", permanent_address="House 14\nLahore")])
        rows = parsed(text)
        assert any(cell == "House 14\nLahore" for cell in rows[1])

    def test_lines_end_the_way_the_standard_says(self) -> None:
        text = csv_text(plan(), [row(serial_no="1")], include_bom=False)
        assert text.count("\r\n") >= 2

    def test_it_is_produced_a_row_at_a_time(self) -> None:
        # A batch of thousands must not be assembled in memory before it is sent.
        columns = plan()
        chunks = list(write_csv(columns, [row(serial_no=str(index)) for index in range(5)]))
        assert len(chunks) == 6  # the header, then one per row


class TestJson:
    def test_each_row_is_an_object_keyed_by_column(self) -> None:
        columns = plan()
        payload = json.loads("".join(write_json(columns, [row(serial_no="1", sex="F")])))
        assert payload[0]["serial_no"] == "1"
        assert payload[0]["sex"] == "F"

    def test_an_empty_export_is_still_valid_json(self) -> None:
        assert json.loads("".join(write_json(plan(), []))) == []

    def test_urdu_is_written_as_itself_not_as_escapes(self) -> None:
        text = "".join(write_json(plan(), [row(serial_no="1", child_full_name="عائشہ")]))
        assert "عائشہ" in text


class TestXlsx:
    def test_the_header_stays_visible_while_scrolling(self) -> None:
        payload = write_xlsx(plan(), [row(serial_no="1")])
        sheet = load_workbook(io.BytesIO(payload)).active
        assert sheet is not None
        assert sheet.freeze_panes == "A2"

    def test_the_header_is_bold_and_filterable(self) -> None:
        payload = write_xlsx(plan(), [row(serial_no="1")])
        sheet = load_workbook(io.BytesIO(payload)).active
        assert sheet is not None
        assert sheet["A1"].font.bold is True
        assert sheet.auto_filter.ref is not None

    def test_columns_are_wide_enough_to_read(self) -> None:
        columns = plan()
        payload = write_xlsx(
            columns, [row(serial_no="1", place_of_birth="Aga Khan University Hospital, Karachi")]
        )
        sheet = load_workbook(io.BytesIO(payload)).active
        assert sheet is not None
        assert all(dimension.width >= 10 for dimension in sheet.column_dimensions.values())

    def test_a_flagged_row_is_tinted(self) -> None:
        columns = plan()
        payload = write_xlsx(
            columns,
            [
                row(serial_no="1", flags=""),
                ExportRow(cells={"serial_no": "2", "flags": "CNIC_INVALID"}, flagged=True),
            ],
        )
        sheet = load_workbook(io.BytesIO(payload)).active
        assert sheet is not None
        clean = sheet.cell(row=2, column=1).fill.fgColor.rgb
        flagged = sheet.cell(row=3, column=1).fill.fgColor.rgb
        assert clean != flagged

    def test_the_flag_is_also_written_as_text(self) -> None:
        # Colour alone is no use to a colour-blind reviewer or to a filter.
        columns = plan()
        payload = write_xlsx(
            columns, [ExportRow(cells={"serial_no": "1", "flags": "CNIC_INVALID"}, flagged=True)]
        )
        sheet = load_workbook(io.BytesIO(payload)).active
        assert sheet is not None
        values = [cell.value for cell in sheet[2]]
        assert "CNIC_INVALID" in values

    def test_urdu_survives_the_workbook(self) -> None:
        columns = plan()
        payload = write_xlsx(columns, [row(serial_no="1", child_full_name="عائشہ نور ملک")])
        sheet = load_workbook(io.BytesIO(payload)).active
        assert sheet is not None
        assert "عائشہ نور ملک" in [cell.value for cell in sheet[2]]

    def test_a_formula_does_not_survive_as_a_formula(self) -> None:
        columns = plan()
        payload = write_xlsx(columns, [row(serial_no="1", certificate_number="=1+1")])
        sheet = load_workbook(io.BytesIO(payload)).active
        assert sheet is not None
        assert "'=1+1" in [cell.value for cell in sheet[2]]
