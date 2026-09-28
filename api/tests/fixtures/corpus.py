"""The rest of the fixture corpus: scans, a bilingual certificate, tables, legacy .doc.

Section 13 of the specification lists the documents the pipeline must be proven on.
The clean text PDF, multi-certificate PDF, .docx, encrypted and corrupt PDFs come
from :mod:`tests.fixtures.builders`; this module adds the rest.

Scans are made the way a scanner makes them: the certificate is rendered to pixels,
then degraded - skewed, rotated, speckled, JPEG-compressed - and the image alone is
wrapped in a PDF with no text layer. OCR therefore meets real artefacts, not a clean
render.
"""

from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass
from pathlib import Path

import fitz
import numpy as np
from docx import Document as DocxDocument
from docx.enum.text import WD_BREAK
from docx.oxml import parse_xml
from PIL import Image, ImageFilter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from tests.fixtures.builders import SAMPLES_BY_KEY, CertificateSample, build_text_pdf

__all__ = [
    "LEGACY_DOC",
    "URDU_BIRTH",
    "arabic_font",
    "build_bilingual_pdf",
    "build_docx_broken",
    "build_docx_content_control",
    "build_docx_line_break",
    "build_docx_merged_cells",
    "build_docx_multi_page",
    "build_docx_textbox",
    "build_docx_urdu",
    "build_docx_with_header",
    "build_scanned_bilingual_pdf",
    "build_scanned_image",
    "build_scanned_pdf",
    "build_table_pdf",
]

LEGACY_DOC = Path(__file__).parent / "files" / "legacy_birth.doc"
"""Binary Word 97-2003 copy of the ``birth_lahore`` sample. See files/README.md."""

_FONT_CANDIDATES = (
    Path(r"C:\Windows\Fonts\arial.ttf"),
    Path(r"C:\Windows\Fonts\tahoma.ttf"),
    Path("/usr/share/fonts/truetype/noto/NotoNaskhArabic-Regular.ttf"),
    Path("/usr/share/fonts/truetype/noto/NotoSansArabic-Regular.ttf"),
    Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
)


def arabic_font() -> Path | None:
    """A font on this machine that has Arabic-script glyphs, if any."""
    return next((path for path in _FONT_CANDIDATES if path.exists()), None)


@dataclass(frozen=True)
class BilingualSample:
    """A certificate with English labels and Urdu values, as issued in Punjab."""

    heading: str
    lines: list[tuple[str, str]]
    expected: dict[str, str]


URDU_BIRTH = BilingualSample(
    heading="CERTIFICATE OF BIRTH / پیدائش کا سرٹیفکیٹ",
    lines=[
        ("Certificate No. / سرٹیفکیٹ نمبر", "BC-2020-001122"),
        ("Name of Child / بچے کا نام", "عائشہ نور ملک"),
        ("Sex / جنس", "لڑکی"),
        ("Date of Birth / تاریخ پیدائش", "14-03-2019"),
        ("Place of Birth / جائے پیدائش", "سروسز ہسپتال، لاہور"),
        ("Father's Name / والد کا نام", "طارق محمود ملک"),
        ("Father's CNIC / والد کا شناختی کارڈ نمبر", "35201-1234567-1"),
        ("Mother's Name / والدہ کا نام", "نسرین اختر"),
        ("Issuing Authority / جاری کنندہ", "یونین کونسل 42، لاہور"),
    ],
    expected={
        "certificate_number": "BC-2020-001122",
        "child_full_name": "عائشہ نور ملک",
        "sex": "F",
        "date_of_birth": "2019-03-14",
        "place_of_birth": "سروسز ہسپتال، لاہور",
        "father_full_name": "طارق محمود ملک",
        "father_id_number": "35201-1234567-1",
        "mother_full_name": "نسرین اختر",
        "issuing_authority": "یونین کونسل 42، لاہور",
    },
)


def build_bilingual_pdf(path: Path, sample: BilingualSample = URDU_BIRTH) -> Path:
    """A text-layer PDF with Urdu shaped and ordered correctly, laid out as a form.

    Every label and value is placed in its own fixed box - left column labels, right
    column values - the way a printed certificate positions its fields. MuPDF's HTML
    layout does the Arabic shaping and bidi ordering inside each box.

    Raises ``RuntimeError`` when no Arabic-capable font is installed; callers skip.
    """
    font = arabic_font()
    if font is None:
        raise RuntimeError("No Arabic-script font is installed on this machine.")

    css = (
        f"@font-face {{font-family: cert; src: url({font.name});}} "
        "* {font-family: cert; font-size: 12px; margin: 0; padding: 0;}"
    )
    archive = fitz.Archive(str(font.parent))
    width, _height = A4
    document = fitz.open()
    page = document.new_page(width=A4[0], height=A4[1])

    page.insert_htmlbox(
        fitz.Rect(50, 50, width - 50, 90),
        f'<p style="text-align:center;font-size:15px">{sample.heading}</p>',
        css=css,
        archive=archive,
    )
    page.insert_htmlbox(
        fitz.Rect(50, 92, width - 50, 110),
        '<p style="text-align:center;font-size:8px">GOVERNMENT OF THE PUNJAB - LOCAL GOVERNMENT</p>',
        css=css,
        archive=archive,
    )
    top = 130.0
    for label, value in sample.lines:
        page.insert_htmlbox(
            fitz.Rect(50, top, 295, top + 22), f"<p>{label}</p>", css=css, archive=archive
        )
        page.insert_htmlbox(
            fitz.Rect(305, top, width - 50, top + 22),
            f'<p dir="auto">{value}</p>',
            css=css,
            archive=archive,
        )
        top += 26.0

    document.save(str(path))
    document.close()
    return path


def _render_page(source: Path, dpi: int) -> Image.Image:
    """The first page of a PDF as grayscale pixels, the way a scanner would see it."""
    buffer = io.BytesIO()
    with fitz.open(str(source)) as document:
        pixmap = document.load_page(0).get_pixmap(dpi=dpi, colorspace=fitz.csGRAY)
        buffer.write(pixmap.tobytes("png"))
    buffer.seek(0)
    with Image.open(buffer) as image:
        return image.convert("L")


def _render_certificate(sample_key: str, dpi: int) -> Image.Image:
    staging = Path(__file__).parent / ".render-cache"
    staging.mkdir(exist_ok=True)
    source = staging / f"{sample_key}.pdf"
    if not source.exists():
        build_text_pdf(source, sample_key)
    return _render_page(source, dpi)


def _degrade(
    image: Image.Image,
    *,
    skew_degrees: float,
    rotation: int,
    noise: float,
    blur: float,
    seed: int,
) -> Image.Image:
    result = image
    if blur > 0:
        result = result.filter(ImageFilter.GaussianBlur(radius=blur))
    if noise > 0:
        rng = np.random.default_rng(seed)
        pixels = np.asarray(result, dtype=np.float32)
        pixels += rng.normal(0.0, 255.0 * noise, size=pixels.shape)
        speckle = rng.random(pixels.shape)
        pixels[speckle < noise / 20] = 0.0
        pixels[speckle > 1 - noise / 20] = 255.0
        result = Image.fromarray(np.clip(pixels, 0, 255).astype(np.uint8), mode="L")
    if skew_degrees:
        result = result.rotate(
            skew_degrees, resample=Image.Resampling.BICUBIC, expand=True, fillcolor=255
        )
    if rotation:
        result = result.rotate(-rotation, expand=True, fillcolor=255)
    return result


def build_scanned_image(
    path: Path,
    *,
    sample_key: str = "birth_lahore",
    dpi: int = 300,
    skew_degrees: float = 0.0,
    rotation: int = 0,
    noise: float = 0.0,
    blur: float = 0.0,
    seed: int = 11,
) -> Path:
    """A degraded scan of a certificate as a standalone image (format from suffix)."""
    image = _degrade(
        _render_certificate(sample_key, dpi),
        skew_degrees=skew_degrees,
        rotation=rotation,
        noise=noise,
        blur=blur,
        seed=seed,
    )
    save_kwargs: dict[str, object] = {"dpi": (dpi, dpi)}
    if path.suffix.lower() in (".jpg", ".jpeg"):
        save_kwargs["quality"] = 85
    image.save(path, **save_kwargs)
    return path


def build_scanned_pdf(
    path: Path,
    *,
    sample_key: str = "birth_lahore",
    dpi: int = 300,
    skew_degrees: float = 0.0,
    rotation: int = 0,
    noise: float = 0.0,
    blur: float = 0.0,
    seed: int = 11,
) -> Path:
    """An image-only PDF - no text layer at all - as a flatbed scanner produces.

    The page image is JPEG-compressed like most scanner output, and the PDF page is
    sized so the image lands at ``dpi``.
    """
    image = _degrade(
        _render_certificate(sample_key, dpi),
        skew_degrees=skew_degrees,
        rotation=rotation,
        noise=noise,
        blur=blur,
        seed=seed,
    )
    jpeg = io.BytesIO()
    image.save(jpeg, format="JPEG", quality=85)

    width_pt = image.width * 72.0 / dpi
    height_pt = image.height * 72.0 / dpi
    document = fitz.open()
    page = document.new_page(width=width_pt, height=height_pt)
    page.insert_image(page.rect, stream=jpeg.getvalue())
    document.save(str(path), deflate=True)
    document.close()
    return path


def build_scanned_bilingual_pdf(
    path: Path,
    *,
    dpi: int = 300,
    skew_degrees: float = 0.0,
    rotation: int = 0,
    noise: float = 0.0,
    blur: float = 0.0,
    seed: int = 11,
) -> Path:
    """An image-only scan of the bilingual certificate: the Urdu OCR case.

    Raises ``RuntimeError`` when no Arabic-script font is installed; callers skip.
    """
    staging = Path(__file__).parent / ".render-cache"
    staging.mkdir(exist_ok=True)
    source = staging / "bilingual.pdf"
    if not source.exists():
        build_bilingual_pdf(source)

    image = _degrade(
        _render_page(source, dpi),
        skew_degrees=skew_degrees,
        rotation=rotation,
        noise=noise,
        blur=blur,
        seed=seed,
    )
    jpeg = io.BytesIO()
    image.save(jpeg, format="JPEG", quality=85)

    document = fitz.open()
    page = document.new_page(width=image.width * 72.0 / dpi, height=image.height * 72.0 / dpi)
    page.insert_image(page.rect, stream=jpeg.getvalue())
    document.save(str(path), deflate=True)
    document.close()
    return path


def build_table_pdf(path: Path, sample_key: str = "death_karachi") -> Path:
    """A certificate laid out as a ruled table, so table detection has lines to find."""
    sample: CertificateSample = SAMPLES_BY_KEY[sample_key]
    styles = getSampleStyleSheet()
    data = [[label, value] for label, value in sample.lines]
    table = Table(data, colWidths=[70 * mm, 100 * mm])
    table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.6, colors.black),
                ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 10),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )
    document = SimpleDocTemplate(str(path), pagesize=A4, title=sample.heading)
    document.build([Paragraph(sample.heading, styles["Title"]), Spacer(1, 8 * mm), table])
    return path


# ---------------------------------------------------------------------------
# Word document shapes
#
# A .docx can hold a certificate in more ways than a table: behind a page break,
# inside merged cells, in a floating text box, or inside a content control - the
# fillable field of a Word form template. Each shape below is one of those, built
# as Word builds it. The raw XML fixtures are minimal Word-shaped skeletons: they
# carry the element structure the reader navigates, not every attribute Word
# writes.
# ---------------------------------------------------------------------------
_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_V_NS = "urn:schemas-microsoft-com:vml"
_MC_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"
_WPS_NS = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
_WP_NS = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
_A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def build_docx_with_header(path: Path, sample_key: str = "birth_lahore") -> Path:
    """A certificate with a page header and footer, as an office template has."""
    sample = SAMPLES_BY_KEY[sample_key]
    document = DocxDocument()
    section = document.sections[0]
    section.header.paragraphs[0].text = "GOVERNMENT OF THE PUNJAB"
    section.footer.paragraphs[0].text = "Verify at punjab.gov.pk"
    document.add_heading(sample.heading, level=1)
    for label, value in sample.lines:
        document.add_paragraph(f"{label}: {value}")
    document.save(str(path))
    return path


def build_docx_multi_page(path: Path, sample_keys: list[str] | None = None) -> Path:
    """Several certificates in one file, separated by explicit page breaks.

    The first break is a break character, the second is "page break before" on the
    following heading - Word writes both, and each must start a page.
    """
    keys = sample_keys or ["birth_lahore", "death_karachi", "marriage_islamabad"]
    document = DocxDocument()
    for index, key in enumerate(keys):
        sample = SAMPLES_BY_KEY[key]
        heading = document.add_paragraph(sample.heading)
        if index == 1:
            document.paragraphs[-2].add_run().add_break(WD_BREAK.PAGE)
        elif index == 2:
            heading.paragraph_format.page_break_before = True
        for label, value in sample.lines[:4]:
            document.add_paragraph(f"{label}: {value}")
    document.save(str(path))
    return path


def build_docx_merged_cells(path: Path) -> Path:
    """A table whose header row spans both columns, and a label merged down a column."""
    document = DocxDocument()
    table = document.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).merge(table.cell(0, 1)).text = "CERTIFICATE OF BIRTH"
    table.cell(1, 0).text = "Name of Child"
    table.cell(1, 1).text = "Ayesha Noor Malik"
    table.cell(2, 0).text = "Permanent Address"
    table.cell(2, 1).text = "House 14, Street 7"
    table.cell(3, 1).text = "Gulberg III, Lahore"
    table.cell(2, 0).merge(table.cell(3, 0))
    document.save(str(path))
    return path


def build_docx_textbox(path: Path, *, with_fallback: bool = False) -> Path:
    """A certificate with a floating text box holding the office seal caption.

    ``with_fallback`` writes the box the way modern Word does: a DrawingML shape
    plus a VML copy inside ``mc:Fallback`` for readers that cannot handle it. The
    text must be read once, not twice.
    """
    document = DocxDocument()
    document.add_paragraph("CERTIFICATE OF BIRTH")
    anchor = document.add_paragraph("Certificate No.: BC-2019-004471")

    vml_box = (
        f'<w:pict xmlns:w="{_W_NS}" xmlns:v="{_V_NS}">'
        "<v:shape><v:textbox><w:txbxContent>"
        "<w:p><w:r><w:t>OFFICE SEAL</w:t></w:r></w:p>"
        "</w:txbxContent></v:textbox></v:shape></w:pict>"
    )
    if with_fallback:
        markup = (
            f'<w:p xmlns:w="{_W_NS}" xmlns:mc="{_MC_NS}" xmlns:wps="{_WPS_NS}" '
            f'xmlns:wp="{_WP_NS}" xmlns:a="{_A_NS}" xmlns:v="{_V_NS}">'
            "<w:r><mc:AlternateContent>"
            '<mc:Choice Requires="wps"><w:drawing><wp:inline><a:graphic><a:graphicData>'
            "<wps:wsp><wps:txbx><w:txbxContent>"
            "<w:p><w:r><w:t>OFFICE SEAL</w:t></w:r></w:p>"
            "</w:txbxContent></wps:txbx></wps:wsp>"
            "</a:graphicData></a:graphic></wp:inline></w:drawing></mc:Choice>"
            f"<mc:Fallback>{vml_box}</mc:Fallback>"
            "</mc:AlternateContent></w:r></w:p>"
        )
    else:
        markup = f'<w:p xmlns:w="{_W_NS}"><w:r>{vml_box}</w:r></w:p>'

    anchor._p.addnext(parse_xml(markup))
    document.save(str(path))
    return path


def build_docx_content_control(path: Path) -> Path:
    """A form template: the registrar's name sits inside a content control."""
    document = DocxDocument()
    document.add_paragraph("CERTIFICATE OF BIRTH")
    anchor = document.add_paragraph("Certificate No.: BC-2019-004471")
    markup = (
        f'<w:sdt xmlns:w="{_W_NS}"><w:sdtPr/><w:sdtContent>'
        "<w:p><w:r><w:t>Registrar: Muhammad Aslam</w:t></w:r></w:p>"
        "</w:sdtContent></w:sdt>"
    )
    anchor._p.addnext(parse_xml(markup))
    document.save(str(path))
    return path


def build_docx_urdu(path: Path, sample: BilingualSample = URDU_BIRTH) -> Path:
    """The bilingual certificate as a Word table, with Urdu in the value column."""
    document = DocxDocument()
    document.add_paragraph(sample.heading)
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in sample.lines:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
    document.save(str(path))
    return path


def build_docx_line_break(path: Path) -> Path:
    """An address split by a manual line break inside one paragraph."""
    document = DocxDocument()
    paragraph = document.add_paragraph("Permanent Address: House 14, Street 7")
    paragraph.runs[0].add_break(WD_BREAK.LINE)
    paragraph.add_run("Gulberg III, Lahore")
    document.save(str(path))
    return path


def build_docx_broken(path: Path, *, shape: str = "garbage") -> Path:
    """A file that claims to be a .docx and is not.

    ``garbage`` is a zip header followed by rubbish - a truncated download.
    ``no_main_part`` is a valid zip whose package has no main document part, which
    is what a partially written export looks like.
    """
    if shape == "garbage":
        path.write_bytes(b"PK" + b"not a word document" * 8)
        return path
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr(
            "[Content_Types].xml",
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="xml" ContentType="application/xml"/></Types>',
        )
    return path
