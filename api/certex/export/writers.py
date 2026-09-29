"""Writing the sheet: CSV, XLSX and JSON.

Three decisions here are about the file being opened by a person on a Windows machine,
in Excel, in Pakistan.

**A leading apostrophe on anything that could be a formula.** A cell beginning ``=``,
``+``, ``-`` or ``@`` is executed by Excel when the file is opened. A certificate number
of "-2019-004471" is not an attack, but a malicious document is a real way in, and the
office that opens the file is the one that pays. The apostrophe is Excel's own escape: it
displays the text and executes nothing.

**A byte-order mark by default.** Without it, Excel on Windows reads a UTF-8 CSV as the
system code page and every Urdu name becomes mojibake. It is optional, because anything
that is not Excel would rather not have it.

**Streaming, not building.** A batch can be thousands of rows and the register can hold
millions; the CSV is produced a row at a time so memory is flat and the browser starts
downloading immediately. Each text format has an asynchronous twin, for rows that arrive
from the database a page at a time - a synchronous writer forces its caller to collect
everything first, which is the same as not streaming at all. XLSX cannot stream either
way - openpyxl's write-only mode still assembles a zip - so it is built in memory, which
is the right trade for a format people open by hand rather than pipe.
"""

from __future__ import annotations

import csv
import io
import json
from collections.abc import AsyncIterable, AsyncIterator, Iterable, Iterator
from typing import Any, Final

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from certex.export.columns import ColumnPlan
from certex.export.rows import ExportRow

__all__ = [
    "DELIMITERS",
    "UTF8_BOM",
    "guard_formula",
    "write_csv",
    "write_csv_stream",
    "write_json",
    "write_json_stream",
    "write_xlsx",
]

UTF8_BOM: Final = "﻿"

DELIMITERS: Final[dict[str, str]] = {
    "comma": ",",
    "semicolon": ";",
    "tab": "\t",
    "pipe": "|",
}
"""Named rather than free-form: a delimiter is a choice from a list, not arbitrary input."""

_FORMULA_STARTERS: Final = ("=", "+", "-", "@", "\t", "\r")
_MAX_COLUMN_WIDTH: Final = 60
_MIN_COLUMN_WIDTH: Final = 10

_FLAGGED_FILL: Final = PatternFill("solid", fgColor="FFF4E5")
"""A warm tint on rows carrying a validation flag - alongside the Flags column, never
instead of it, because colour alone is no use to a colour-blind reviewer or a filter."""


def guard_formula(value: str) -> str:
    """Neutralise a cell a spreadsheet would otherwise execute."""
    if value and value.startswith(_FORMULA_STARTERS):
        return f"'{value}"
    return value


class _CsvBuffer:
    """One row of CSV at a time, written through the standard library.

    Exists so the synchronous and asynchronous writers below format rows through
    exactly the same code - RFC 4180 quoting, CRLF endings, formula guarding - and
    cannot drift into two dialects of the same file format.
    """

    def __init__(self, *, delimiter: str) -> None:
        self._buffer = io.StringIO()
        self._writer = csv.writer(
            self._buffer, delimiter=delimiter, quoting=csv.QUOTE_MINIMAL, lineterminator="\r\n"
        )

    def _flush(self) -> str:
        text = self._buffer.getvalue()
        self._buffer.seek(0)
        self._buffer.truncate(0)
        return text

    def header(self, plan: ColumnPlan, *, include_bom: bool) -> str:
        self._writer.writerow([guard_formula(header) for header in plan.headers])
        return (UTF8_BOM if include_bom else "") + self._flush()

    def row(self, plan: ColumnPlan, row: ExportRow) -> str:
        self._writer.writerow([guard_formula(value) for value in row.values(plan)])
        return self._flush()


def write_csv(
    plan: ColumnPlan,
    rows: Iterable[ExportRow],
    *,
    delimiter: str = ",",
    include_bom: bool = True,
) -> Iterator[str]:
    """Yield the CSV a piece at a time: the header, then one chunk per row.

    RFC 4180 quoting throughout, and CRLF line endings, which is what the standard says
    and what Excel expects.
    """
    buffer = _CsvBuffer(delimiter=delimiter)
    yield buffer.header(plan, include_bom=include_bom)
    for row in rows:
        yield buffer.row(plan, row)


async def write_csv_stream(
    plan: ColumnPlan,
    rows: AsyncIterable[ExportRow],
    *,
    delimiter: str = ",",
    include_bom: bool = True,
) -> AsyncIterator[str]:
    """The same CSV, from rows that arrive a page at a time from the database.

    The point of the asynchronous twin is that nothing is ever wholly in memory. The
    register holds millions of entries, and collecting them into a list to hand to the
    synchronous writer would undo the streaming it was written for.
    """
    buffer = _CsvBuffer(delimiter=delimiter)
    yield buffer.header(plan, include_bom=include_bom)
    async for row in rows:
        yield buffer.row(plan, row)


def _json_object(plan: ColumnPlan, row: ExportRow, *, first: bool) -> str:
    payload = {key: row.cells.get(key, "") for key in plan.keys}
    prefix = "" if first else ",\n"
    return prefix + "  " + json.dumps(payload, ensure_ascii=False)


def write_json(plan: ColumnPlan, rows: Iterable[ExportRow]) -> Iterator[str]:
    """Yield a JSON array, one object per row, streamed like the CSV.

    Objects rather than arrays-of-arrays: whatever reads this next should not have to
    know the column order to make sense of it.
    """
    yield "[\n"
    first = True
    for row in rows:
        yield _json_object(plan, row, first=first)
        first = False
    yield "\n]\n"


async def write_json_stream(plan: ColumnPlan, rows: AsyncIterable[ExportRow]) -> AsyncIterator[str]:
    """The same JSON array, from rows arriving a page at a time."""
    yield "[\n"
    first = True
    async for row in rows:
        yield _json_object(plan, row, first=first)
        first = False
    yield "\n]\n"


def write_xlsx(plan: ColumnPlan, rows: Iterable[ExportRow]) -> bytes:
    """Build the workbook: frozen header, sized columns, flagged rows tinted."""
    workbook = Workbook()
    sheet = workbook.active
    if sheet is None:  # pragma: no cover - a new workbook always has one
        sheet = workbook.create_sheet()
    sheet.title = "Certificates"

    sheet.append([guard_formula(header) for header in plan.headers])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    # The header stays visible while a clerk scrolls a thousand rows.
    sheet.freeze_panes = "A2"

    widths = [len(header) for header in plan.headers]
    for row in rows:
        values = [guard_formula(value) for value in row.values(plan)]
        sheet.append(values)
        for index, value in enumerate(values):
            widths[index] = max(widths[index], min(len(value), _MAX_COLUMN_WIDTH))
        if row.flagged:
            for cell in sheet[sheet.max_row]:
                cell.fill = _FLAGGED_FILL

    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(
            _MIN_COLUMN_WIDTH, min(width + 2, _MAX_COLUMN_WIDTH)
        )
    # Excel's own filter row, so the first thing a clerk wants to do already works.
    sheet.auto_filter.ref = sheet.dimensions

    payload = io.BytesIO()
    workbook.save(payload)
    return payload.getvalue()


def json_rows(plan: ColumnPlan, rows: Iterable[ExportRow]) -> list[dict[str, Any]]:
    """The same objects ``write_json`` streams, for callers that want them in memory."""
    return [{key: row.cells.get(key, "") for key in plan.keys} for row in rows]
