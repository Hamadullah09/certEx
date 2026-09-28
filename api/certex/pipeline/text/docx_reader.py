"""Word (.docx) documents.

A .docx carries structure but no geometry: there is no page layout until Word
renders it. Everything the rules engine needs is still here - reading order,
tables, which cell a value sits in - so this reader produces a layout with
*synthetic* positions:

* lines are stacked top to bottom in document order;
* a table row becomes one line with its cells side by side, each cell spanning its
  column's share of the width, so "the value is to the right of the label" holds
  for a two-column certificate table exactly as it does on a scanned form.

Read, in order: the section header, the body (paragraphs and tables interleaved as
they appear, each paragraph followed by any text boxes anchored in it), and the
section footer. An explicit page break - a break character or "page break before"
- starts a new page, which is how a Word file holding several certificates
separates them; a manual line break starts a new line.

Content controls (the fillable fields of a Word form template) are read through, at
body, row, cell and run level. Where Word stores alternate content, the primary
choice is read and the ``mc:Fallback`` copy for older readers is ignored.
"""

from __future__ import annotations

import zipfile
from collections.abc import Iterator
from dataclasses import dataclass, field
from pathlib import Path
from typing import Final, Protocol

from docx import Document as DocxDocument
from docx.opc.exceptions import OpcError
from docx.oxml.ns import qn

from certex.core.errors import CorruptDocumentError
from certex.pipeline.text.normalize import normalize_text
from certex.schemas.layout import BBox, Block, Line, PageLayout, Table, TableCell, Word

__all__ = ["read_docx_pages"]

_MIN_SLOTS = 40
"""Pages are laid out on at least this many line slots, so a short certificate's
lines keep a realistic height instead of each filling a tenth of the page."""

_W_P: Final = qn("w:p")
_W_TBL: Final = qn("w:tbl")
_W_TR: Final = qn("w:tr")
_W_TC: Final = qn("w:tc")
_W_T: Final = qn("w:t")
_W_TAB: Final = qn("w:tab")
_W_BR: Final = qn("w:br")
_W_CR: Final = qn("w:cr")
_W_NO_BREAK_HYPHEN: Final = qn("w:noBreakHyphen")
_W_TYPE: Final = qn("w:type")
_W_TXBX_CONTENT: Final = qn("w:txbxContent")
_W_SDT: Final = qn("w:sdt")
_W_SDT_CONTENT: Final = qn("w:sdtContent")
_W_CUSTOM_XML: Final = qn("w:customXml")
_W_TC_PR: Final = qn("w:tcPr")
_W_H_MERGE: Final = qn("w:hMerge")
_W_VAL: Final = qn("w:val")
_MC_FALLBACK: Final = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"


class _Element(Protocol):
    """The slice of lxml's element API this reader uses (python-docx ships no types)."""

    @property
    def tag(self) -> str: ...

    @property
    def text(self) -> str | None: ...

    def iterchildren(self) -> Iterator[_Element]: ...

    def iterancestors(self) -> Iterator[_Element]: ...

    def iter(self, *tags: str) -> Iterator[_Element]: ...

    def find(self, path: str) -> _Element | None: ...

    def get(self, key: str) -> str | None: ...


@dataclass(slots=True)
class _Row:
    """One synthetic line: plain paragraph text, or the cells of a table row."""

    cells: list[str]
    is_table_row: bool = False


@dataclass(slots=True)
class _DraftPage:
    rows: list[_Row] = field(default_factory=list)
    tables: list[list[list[str]]] = field(default_factory=list)
    table_row_offsets: list[int] = field(default_factory=list)


def _inside(node: _Element, ancestor_tag: str, *, stop: _Element) -> bool:
    """Whether ``node`` has an ancestor with ``ancestor_tag`` below ``stop``."""
    for ancestor in node.iterancestors():
        if ancestor is stop:
            return False
        if ancestor.tag == ancestor_tag:
            return True
    return False


def _children(element: _Element, tags: tuple[str, ...]) -> Iterator[_Element]:
    """Direct children with the given tags, looking through content controls."""
    for child in element.iterchildren():
        if child.tag in tags:
            yield child
        elif child.tag == _W_SDT:
            content = child.find(_W_SDT_CONTENT)
            if content is not None:
                yield from _children(content, tags)
        elif child.tag == _W_CUSTOM_XML:
            yield from _children(child, tags)


def _paragraph_pages(paragraph: _Element) -> list[list[str]]:
    """A paragraph's text lines, grouped by the explicit page breaks inside it.

    Text inside text boxes anchored in the paragraph is excluded; it is read
    separately by :func:`_textboxes`.
    """
    pages: list[list[str]] = [[]]
    pieces: list[str] = []

    def end_line() -> None:
        text = normalize_text("".join(pieces))
        if text:
            pages[-1].append(text)
        pieces.clear()

    for node in paragraph.iter(_W_T, _W_TAB, _W_BR, _W_CR, _W_NO_BREAK_HYPHEN):
        if _inside(node, _W_TXBX_CONTENT, stop=paragraph) or _inside(
            node, _MC_FALLBACK, stop=paragraph
        ):
            continue
        if node.tag == _W_T:
            pieces.append(node.text or "")
        elif node.tag == _W_TAB:
            pieces.append("\t")
        elif node.tag == _W_NO_BREAK_HYPHEN:
            pieces.append("-")
        elif node.tag == _W_BR and node.get(_W_TYPE) == "page":
            end_line()
            pages.append([])
        else:
            # A manual line break, a column break or a carriage return.
            end_line()
    end_line()
    return pages


def _textboxes(paragraph: _Element) -> list[str]:
    """Lines of the text boxes anchored in a paragraph.

    Word writes a modern text box twice - the DrawingML shape, and a VML copy inside
    ``mc:Fallback`` for older readers - so the fallback copy is skipped.
    """
    lines: list[str] = []
    for box in paragraph.iter(_W_TXBX_CONTENT):
        if _inside(box, _MC_FALLBACK, stop=paragraph) or _inside(
            box, _W_TXBX_CONTENT, stop=paragraph
        ):
            # The fallback copy, or a text box nested in another (read with its parent).
            continue
        for inner in box.iter(_W_P):
            if _inside(inner, _W_TXBX_CONTENT, stop=box):
                continue  # read below, through the paragraph that anchors it
            for page in _paragraph_pages(inner):
                lines.extend(page)
            lines.extend(_textboxes(inner))
    return lines


def _cell_text(cell: _Element) -> str:
    """A table cell's text on one line, including nested tables' text."""
    lines: list[str] = []
    for paragraph in cell.iter(_W_P):
        if _inside(paragraph, _W_TXBX_CONTENT, stop=cell) or _inside(
            paragraph, _MC_FALLBACK, stop=cell
        ):
            continue
        for page in _paragraph_pages(paragraph):
            lines.extend(page)
        lines.extend(_textboxes(paragraph))
    return " ".join(lines)


def _continues_horizontal_merge(cell: _Element) -> bool:
    """A legacy ``hMerge`` continuation cell, whose text lives in the cell before it."""
    merge = cell.find(f"{_W_TC_PR}/{_W_H_MERGE}")
    return merge is not None and merge.get(_W_VAL) != "restart"


def _row_cells(row: _Element) -> list[str]:
    """Row cell texts, with a cell spanning several grid columns counted once.

    A modern spanning cell is one ``w:tc`` with a grid span; a legacy one is a run
    of cells joined by ``hMerge``, and its continuation cells are skipped.
    """
    return [
        _cell_text(cell)
        for cell in _children(row, (_W_TC,))
        if not _continues_horizontal_merge(cell)
    ]


def _page_break_before(paragraph: _Element) -> bool:
    return paragraph.find(f"{qn('w:pPr')}/{qn('w:pageBreakBefore')}") is not None


def _body_items(document: object) -> Iterator[_Element]:
    body: _Element = document.element.body  # type: ignore[attr-defined]
    return _children(body, (_W_P, _W_TBL))


def _header_footer_texts(document: object) -> tuple[list[str], list[str]]:
    headers: list[str] = []
    footers: list[str] = []
    for section in document.sections:  # type: ignore[attr-defined]
        for part, target in ((section.header, headers), (section.footer, footers)):
            if part.is_linked_to_previous and target:
                continue
            part_element: _Element = part._element
            for child in _children(part_element, (_W_P, _W_TBL)):
                texts: list[str]
                if child.tag == _W_P:
                    texts = [line for page in _paragraph_pages(child) for line in page]
                    texts.extend(_textboxes(child))
                elif child.tag == _W_TBL:
                    texts = [
                        joined
                        for row in _children(child, (_W_TR,))
                        if (joined := " ".join(cell for cell in _row_cells(row) if cell))
                    ]
                else:
                    continue
                target.extend(text for text in texts if text not in target)
    return headers, footers


def _draft_pages(document: object) -> list[_DraftPage]:
    pages = [_DraftPage()]

    for item in _body_items(document):
        if item.tag == _W_P:
            if _page_break_before(item) and pages[-1].rows:
                pages.append(_DraftPage())
            for index, lines in enumerate(_paragraph_pages(item)):
                if index > 0:
                    pages.append(_DraftPage())
                pages[-1].rows.extend(_Row(cells=[line]) for line in lines)
            pages[-1].rows.extend(_Row(cells=[line]) for line in _textboxes(item))
            continue

        current = pages[-1]
        grid: list[list[str]] = []
        current.table_row_offsets.append(len(current.rows))
        for row in _children(item, (_W_TR,)):
            cells = _row_cells(row)
            if any(cells):
                grid.append(cells)
                current.rows.append(_Row(cells=cells, is_table_row=True))
        current.tables.append(grid)

    return [page for page in pages if page.rows] or [_DraftPage()]


def _layout_page(
    draft: _DraftPage, *, page_number: int, headers: list[str], footers: list[str]
) -> PageLayout:
    rows = (
        [_Row(cells=[text]) for text in headers]
        + draft.rows
        + [_Row(cells=[text]) for text in footers]
    )
    header_offset = len(headers)
    slot = 1.0 / max(len(rows), _MIN_SLOTS)

    words: list[Word] = []
    lines: list[Line] = []
    row_boxes: list[list[BBox]] = []

    for row_index, row in enumerate(rows):
        top = row_index * slot
        bottom = top + slot * 0.8
        columns = len(row.cells)
        line_word_indices: list[int] = []
        cell_boxes: list[BBox] = []
        for column, cell_text in enumerate(row.cells):
            left = column / columns
            right = (column + 1) / columns
            cell_boxes.append(BBox(x0=left, y0=top, x1=right, y1=bottom))
            tokens = cell_text.split()
            total = sum(len(token) + 1 for token in tokens) or 1
            cursor = left
            span = right - left
            for token in tokens:
                share = span * (len(token) + 1) / total
                line_word_indices.append(len(words))
                words.append(
                    Word(
                        text=token,
                        bbox=BBox(
                            x0=round(cursor, 5),
                            y0=round(top, 5),
                            x1=round(min(cursor + share * 0.9, right), 5),
                            y1=round(bottom, 5),
                        ),
                        line=len(lines),
                    )
                )
                cursor += share
        row_boxes.append(cell_boxes)
        if not line_word_indices:
            continue
        line_box = BBox.enclosing([words[index].bbox for index in line_word_indices])
        if line_box is None:  # pragma: no cover
            continue
        lines.append(
            Line(
                # Cells joined with a tab-like gap, so a regex over the line text still
                # sees the label and value as separate columns.
                text="   ".join(cell for cell in row.cells if cell),
                bbox=line_box,
                words=line_word_indices,
                block=0,
            )
        )

    tables: list[Table] = []
    for grid, offset in zip(draft.tables, draft.table_row_offsets, strict=True):
        start = header_offset + offset
        table_rows: list[list[TableCell]] = []
        for index, cells in enumerate(grid):
            boxes = row_boxes[start + index] if start + index < len(row_boxes) else []
            table_rows.append(
                [
                    TableCell(text=text, bbox=boxes[column] if column < len(boxes) else None)
                    for column, text in enumerate(cells)
                ]
            )
        if table_rows:
            tables.append(Table(rows=table_rows))

    blocks = (
        [
            Block(
                bbox=BBox.enclosing([line.bbox for line in lines]) or BBox(x0=0, y0=0, x1=1, y1=1),
                lines=list(range(len(lines))),
            )
        ]
        if lines
        else []
    )
    return PageLayout(
        page_number=page_number,
        unit="synthetic",
        engine="docx",
        words=words,
        lines=lines,
        blocks=blocks,
        tables=tables,
    )


_DOCX_ERRORS: tuple[type[BaseException], ...] = (
    OpcError,  # not a Word package at all
    zipfile.BadZipFile,  # a damaged zip container
    KeyError,  # a package missing its main document part
    SyntaxError,  # malformed XML inside the package (lxml's XMLSyntaxError)
    ValueError,
)


def read_docx_pages(path: Path) -> list[PageLayout]:
    """Read a .docx into one layout per page (explicit page breaks split pages).

    Raises :class:`CorruptDocumentError` when the file is not a readable Word document.
    """
    try:
        document = DocxDocument(str(path))
        headers, footers = _header_footer_texts(document)
        drafts = _draft_pages(document)
    except _DOCX_ERRORS as exc:
        raise CorruptDocumentError(
            "The Word document could not be read; it appears damaged.",
            remediation="Open it in Word, save a new copy, and upload that copy.",
        ) from exc
    return [
        _layout_page(draft, page_number=index + 1, headers=headers, footers=footers)
        for index, draft in enumerate(drafts)
    ]
